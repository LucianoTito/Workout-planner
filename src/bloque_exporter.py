"""
Exportador a Excel para BLOQUES.

Reutiliza el ESTILO del exportador CrossFit (temas de color, grilla, alturas)
importándolo de excel_exporter.py, para que ambos programas se vean igual.
No modifica excel_exporter.py: solo consume sus constantes y clase de estilos.

Layout (idéntico en espíritu al CrossFit):
  Fila 1 : título del bloque + semana + fase + atleta (mergeada)
  Fila 2 : nota global de la rodilla (mergeada, si existe)
  Fila 3 : headers de día (DÍA A, DÍA B, ...)
  Fila 4+: subtítulo (tipo de día + duración)
  Luego  : cada día fluye sus bloques (secciones) hacia abajo
  Final  : bloque NOTAS (área de escritura libre)
"""

import openpyxl
from openpyxl.utils import get_column_letter
from pathlib import Path

from src.excel_exporter import (
    TEMAS, TEMA_DEFAULT, _Estilos, BORDER_THIN,
    ALIGN_CENTER, ALIGN_LEFT, filas_necesarias,
    ANCHO_COLUMNA, ALTURA_FILA, CHARS_POR_LINEA,
    FILA_NOTAS, FILAS_AREA_NOTAS,
)
from src.bloque_models import SemanaBloque


def _escribir_celda(ws, fila, col, valor, font, fill, filas, align):
    """Escribe una celda; si necesita >1 fila, combina verticalmente y centra."""
    celda = ws.cell(row=fila, column=col)
    celda.value = valor
    celda.font = font
    if fill:
        celda.fill = fill
    celda.alignment = align
    celda.border = BORDER_THIN

    if filas > 1:
        ws.merge_cells(
            start_row=fila, start_column=col,
            end_row=fila + filas - 1, end_column=col,
        )
        for f in range(fila, fila + filas):
            cc = ws.cell(row=f, column=col)
            cc.border = BORDER_THIN
            if fill:
                cc.fill = fill


class BloqueExporter:
    """Exporta una semana de un bloque a un archivo Excel."""

    def __init__(self, output_dir: str = "output", tema: str = TEMA_DEFAULT):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.tema = tema if tema in TEMAS else TEMA_DEFAULT

    def exportar_semana(self, semana: SemanaBloque,
                        nombre_archivo: str | None = None) -> str:
        if not nombre_archivo:
            nombre_archivo = (
                f"{semana.nombre_bloque.replace(' ', '_')}"
                f"_S{semana.numero_semana}_{semana.atleta}"
            )

        est = _Estilos(TEMAS[self.tema])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Semana {semana.numero_semana}"

        n_dias = len(semana.dias)
        for col in range(1, n_dias + 1):
            ws.column_dimensions[get_column_letter(col)].width = ANCHO_COLUMNA

        fila = 1

        # ─── FILA 1: título del bloque (mergeada) ───
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_dias)
        c = ws.cell(row=fila, column=1)
        c.value = (
            f"{semana.nombre_bloque.upper()} · "
            f"SEMANA {semana.numero_semana} de {semana.total_semanas} — "
            f"{semana.fase} — {semana.atleta}"
        )
        c.font = est.font_semana
        if est.fill_semana:
            c.fill = est.fill_semana
        c.alignment = ALIGN_CENTER
        ws.row_dimensions[fila].height = ALTURA_FILA
        fila += 1

        # ─── FILA 2: nota global (mergeada a lo ancho) ───
        if semana.nota_global:
            nf = filas_necesarias(semana.nota_global, CHARS_POR_LINEA * n_dias)
            ws.merge_cells(start_row=fila, start_column=1,
                           end_row=fila + nf - 1, end_column=n_dias)
            c = ws.cell(row=fila, column=1)
            c.value = semana.nota_global
            c.font = est.font_sub
            c.alignment = ALIGN_CENTER
            for f in range(fila, fila + nf):
                for col in range(1, n_dias + 1):
                    cc = ws.cell(row=f, column=col)
                    cc.border = BORDER_THIN
                    if est.fill_sub:
                        cc.fill = est.fill_sub
                ws.row_dimensions[f].height = ALTURA_FILA
            fila += nf

        # ─── FILA: headers de día ───
        for col, dia in enumerate(semana.dias, 1):
            c = ws.cell(row=fila, column=col)
            c.value = dia.titulo
            c.font = est.font_dia
            c.fill = est.fill_dia
            c.alignment = ALIGN_CENTER
            c.border = BORDER_THIN
        ws.row_dimensions[fila].height = ALTURA_FILA
        fila += 1

        # ─── FILA: subtítulos (tipo de día + duración) ───
        max_filas_sub = max(filas_necesarias(dia.subtitulo) for dia in semana.dias)
        for col, dia in enumerate(semana.dias, 1):
            _escribir_celda(
                ws, fila, col, dia.subtitulo,
                font=est.font_sub, fill=est.fill_sub,
                filas=max_filas_sub, align=ALIGN_CENTER,
            )
        for f in range(fila, fila + max_filas_sub):
            ws.row_dimensions[f].height = ALTURA_FILA
        fila += max_filas_sub

        fila_inicio_bloques = fila

        # ─── Bloques por columna (cada día fluye a su ritmo) ───
        max_fila_usada = fila_inicio_bloques
        for col, dia in enumerate(semana.dias, 1):
            fila_actual = fila_inicio_bloques
            for header, lineas in dia.bloques:
                nf_header = filas_necesarias(header)
                _escribir_celda(
                    ws, fila_actual, col, header,
                    font=est.font_bloque, fill=est.fill_bloque,
                    filas=nf_header, align=ALIGN_CENTER,
                )
                fila_actual += nf_header

                for linea in lineas:
                    nf = filas_necesarias(linea)
                    align = ALIGN_CENTER if nf > 1 else ALIGN_LEFT
                    _escribir_celda(
                        ws, fila_actual, col, linea,
                        font=est.font_normal, fill=est.fill_contenido,
                        filas=nf, align=align,
                    )
                    fila_actual += nf

            max_fila_usada = max(max_fila_usada, fila_actual)

        for f in range(fila_inicio_bloques, max_fila_usada + 1):
            ws.row_dimensions[f].height = ALTURA_FILA

        # ─── BLOQUE NOTAS (área de escritura libre) ───
        fila_notas = max(FILA_NOTAS, max_fila_usada + 2)
        for col in range(1, n_dias + 1):
            _escribir_celda(
                ws, fila_notas, col, "NOTAS",
                font=est.font_bloque, fill=est.fill_bloque,
                filas=1, align=ALIGN_CENTER,
            )
            for f in range(fila_notas + 1, fila_notas + 1 + FILAS_AREA_NOTAS):
                cc = ws.cell(row=f, column=col)
                cc.border = BORDER_THIN
                if est.fill_contenido:
                    cc.fill = est.fill_contenido
                cc.alignment = ALIGN_LEFT
        for f in range(fila_notas, fila_notas + 1 + FILAS_AREA_NOTAS):
            ws.row_dimensions[f].height = ALTURA_FILA

        ruta = self.output_dir / f"{nombre_archivo}.xlsx"
        wb.save(str(ruta))
        return str(ruta)
