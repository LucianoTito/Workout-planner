"""
Exportador a Excel.
Genera archivos .xlsx con formato profesional para compartir por Drive.

Cada día (columna) renderiza SOLO los bloques que efectivamente tiene (bloques
vacíos se omiten). Cada columna fluye a su propio ritmo (no se fuerza alineación
entre días).

Prolijidad de celdas (Enfoque B):
  - Todas las filas tienen la MISMA altura fija.
  - Cada línea de contenido ocupa 1 fila si entra; si el texto es largo, ocupa
    N filas combinadas verticalmente y centradas — igual a como se hace a mano.
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

from src.models import SemanaEntrenamiento, DiaEntrenamiento, TipoDia


# ═══════════════════════════════════════════════════════════
# DIMENSIONES DE LA GRILLA
# ═══════════════════════════════════════════════════════════
ANCHO_COLUMNA = 44          # ancho de cada columna (unidades Excel)
ALTURA_FILA = 18            # altura fija de cada fila (puntos)
CHARS_POR_LINEA = 34        # aprox. caracteres que entran en una línea a ese ancho
FILA_NOTAS = 46             # fila fija donde arranca el bloque NOTAS
FILAS_AREA_NOTAS = 18       # filas de escritura libre debajo del header NOTAS

# ═══════════════════════════════════════════════════════════
# TEMAS DE COLORES (perfiles)
# ═══════════════════════════════════════════════════════════
# Cada tema define los colores de cada rol de celda.
# Roles:
#   semana    → fila título del ciclo (SEMANA X de 8 …)
#   dia       → headers de día (DÍA 1, DÍA 2 …)
#   subtitulo → tipo de día / "qué entrenamiento es"
#   bloque    → headers de sección (CORE, MUSCULACIÓN, WOD …)
#   contenido → celdas de ejercicios
TEMAS = {
    "rosa": {   # actual — Brisa
        "semana":    {"fill": None,     "font": "A64D79", "bold": True,  "size": 11},
        "dia":       {"fill": "C27BA0", "font": "FFFFFF", "bold": True,  "size": 12},
        "subtitulo": {"fill": "A64D79", "font": "FFFFFF", "bold": True,  "size": 10},
        "bloque":    {"fill": "F4D9E4", "font": "000000", "bold": True,  "size": 12},
        "contenido": {"fill": None,     "font": "000000", "bold": False, "size": 12},
    },
    "arena": {  # nuevo — Luciano
        "semana":    {"fill": "2B2B2B", "font": "FFFFFF", "bold": True,  "size": 11},
        "dia":       {"fill": "2B2B2B", "font": "FFFFFF", "bold": True,  "size": 12},
        "subtitulo": {"fill": "F4ECDD", "font": "7A5A1F", "bold": True,  "size": 10},
        "bloque":    {"fill": "B8863B", "font": "FFFFFF", "bold": True,  "size": 12},
        "contenido": {"fill": "FFFFFF", "font": "1A1A1A", "bold": False, "size": 12},
    },
}
TEMA_DEFAULT = "rosa"


class _Estilos:
    """Construye los objetos Font/Fill a partir de un tema."""

    def __init__(self, tema: dict):
        def f(rol):
            r = tema[rol]
            return Font(name="Arial", bold=r["bold"], color=r["font"], size=r["size"])

        def fl(rol):
            c = tema[rol]["fill"]
            return PatternFill("solid", fgColor=c) if c else None

        self.font_semana = f("semana");    self.fill_semana = fl("semana")
        self.font_dia    = f("dia");       self.fill_dia    = fl("dia")
        self.font_sub    = f("subtitulo"); self.fill_sub    = fl("subtitulo")
        self.font_bloque = f("bloque");    self.fill_bloque = fl("bloque")
        self.font_normal = f("contenido"); self.fill_contenido = fl("contenido")
        self.font_notas  = Font(name="Arial", italic=True, color="808080",
                                size=tema["contenido"]["size"])


BORDER_THIN = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def filas_necesarias(texto: str, chars_por_linea: int = CHARS_POR_LINEA) -> int:
    """
    Estima cuántas filas de altura fija necesita un texto.

    Considera saltos de línea explícitos y el wrap por ancho de columna.
    """
    if not texto:
        return 1
    total = 0
    for parte in str(texto).split("\n"):
        largo = len(parte)
        if largo == 0:
            total += 1
        else:
            # ceil(largo / chars_por_linea)
            total += -(-largo // chars_por_linea)
    return max(1, total)


class ExcelExporter:
    """Exporta semanas de entrenamiento a archivos Excel."""

    def __init__(self, output_dir: str = "output", tema: str = TEMA_DEFAULT):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.tema = tema if tema in TEMAS else TEMA_DEFAULT

    def exportar_semana(self, semana: SemanaEntrenamiento,
                        nombre_archivo: str | None = None) -> str:
        """Exporta una semana a un archivo Excel."""
        if not nombre_archivo:
            nombre_archivo = (
                f"S{semana.numero_semana}_{semana.atleta.nombre}"
                f"_{semana.fecha_inicio or 'sin_fecha'}"
            )

        est = _Estilos(TEMAS[self.tema])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"Semana {semana.numero_semana}"

        n_dias = len(semana.dias)
        for col in range(1, n_dias + 1):
            ws.column_dimensions[get_column_letter(col)].width = ANCHO_COLUMNA

        # Bloques de cada día
        bloques_por_dia = [self._bloques_de_dia(dia) for dia in semana.dias]

        # ─── FILA 1: Info del ciclo (mergeada) ───
        fila = 1
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_dias)
        c = ws.cell(row=fila, column=1)
        c.value = (f"SEMANA {semana.numero_semana} de 8 — "
                   f"{semana.fase} — {semana.atleta.nombre}")
        c.font = est.font_semana
        if est.fill_semana:
            c.fill = est.fill_semana
        c.alignment = ALIGN_CENTER
        ws.row_dimensions[fila].height = ALTURA_FILA
        fila += 1

        # ─── FILA 2: Headers de día ───
        for col, dia in enumerate(semana.dias, 1):
            c = ws.cell(row=fila, column=col)
            c.value = dia.titulo
            c.font = est.font_dia
            c.fill = est.fill_dia
            c.alignment = ALIGN_CENTER
            c.border = BORDER_THIN
        ws.row_dimensions[fila].height = ALTURA_FILA
        fila += 1

        # ─── FILA 3: Subtítulos (tipo de día) — puede necesitar 2 filas ───
        # Calculamos las filas que ocupa el subtítulo más largo
        max_filas_sub = max(
            filas_necesarias(dia.subtitulo) for dia in semana.dias
        )
        for col, dia in enumerate(semana.dias, 1):
            self._escribir_celda(
                ws, fila, col, dia.subtitulo,
                font=est.font_sub, fill=est.fill_sub,
                filas=max_filas_sub, align=ALIGN_CENTER,
            )
        for f in range(fila, fila + max_filas_sub):
            ws.row_dimensions[f].height = ALTURA_FILA
        fila += max_filas_sub

        fila_inicio_bloques = fila

        # ─── Renderizar bloques por columna (cada día fluye) ───
        max_fila_usada = fila_inicio_bloques
        for col, bloques in enumerate(bloques_por_dia, 1):
            fila_actual = fila_inicio_bloques
            for header, lineas in bloques:
                # Header del bloque (centrado, altura fija — normalmente 1 fila)
                nf_header = filas_necesarias(header)
                self._escribir_celda(
                    ws, fila_actual, col, header,
                    font=est.font_bloque, fill=est.fill_bloque,
                    filas=nf_header, align=ALIGN_CENTER,
                )
                fila_actual += nf_header

                # Contenido: cada línea ocupa las filas que necesite
                for linea in lineas:
                    es_wod_vacio = (linea == "[COMPLETAR MANUALMENTE]")
                    nf = filas_necesarias(linea)
                    # Alineación: centrado si ocupa varias filas (merge), izq si 1
                    align = ALIGN_CENTER if nf > 1 else ALIGN_LEFT
                    self._escribir_celda(
                        ws, fila_actual, col, linea,
                        font=est.font_notas if es_wod_vacio else est.font_normal,
                        fill=est.fill_contenido, filas=nf, align=align,
                    )
                    fila_actual += nf

            max_fila_usada = max(max_fila_usada, fila_actual)

        # ─── Fijar altura de TODAS las filas de contenido ───
        for f in range(fila_inicio_bloques, max_fila_usada + 1):
            ws.row_dimensions[f].height = ALTURA_FILA

        # ─── BLOQUE NOTAS (fijo en fila 46, área de escritura libre) ───
        fila_notas = max(FILA_NOTAS, max_fila_usada + 2)   # nunca pisa el contenido
        for col in range(1, n_dias + 1):
            self._escribir_celda(
                ws, fila_notas, col, "NOTAS",
                font=est.font_bloque, fill=est.fill_bloque,
                filas=1, align=ALIGN_CENTER,
            )
            for f in range(fila_notas + 1, fila_notas + 1 + FILAS_AREA_NOTAS):
                cc = ws.cell(row=f, column=col)
                cc.border = BORDER_THIN
                if est.fill_contenido:
                    cc.fill = est.fill_contenido
                cc.alignment = ALIGN_LEFT
        for f in range(fila_notas, fila_notas + 1 + FILAS_AREA_NOTAS):
            ws.row_dimensions[f].height = ALTURA_FILA

        ruta = self.output_dir / f"{nombre_archivo}.xlsx"
        wb.save(str(ruta))
        return str(ruta)

    def _escribir_celda(self, ws, fila, col, valor, font, fill, filas, align):
        """
        Escribe una celda. Si necesita más de una fila, combina verticalmente
        las filas y centra el contenido (como el merge manual).
        """
        celda = ws.cell(row=fila, column=col)
        celda.value = valor
        celda.font = font
        if fill:
            celda.fill = fill
        celda.alignment = align
        celda.border = BORDER_THIN

        if filas > 1:
            # Combinar verticalmente esta celda con las de abajo (misma columna)
            ws.merge_cells(
                start_row=fila, start_column=col,
                end_row=fila + filas - 1, end_column=col,
            )
            # Aplicar borde y fill a todas las celdas del merge
            for f in range(fila, fila + filas):
                cc = ws.cell(row=f, column=col)
                cc.border = BORDER_THIN
                if fill:
                    cc.fill = fill

    def _bloques_de_dia(self, dia: DiaEntrenamiento) -> list[tuple[str, list[str]]]:
        """
        Construye la lista de bloques de un día, EN ORDEN, omitiendo los vacíos.
        """
        bloques = []

        # ─── CORE / ESTABILIDAD ───
        cb = dia.core_block
        if cb and cb.ejercicios:
            lineas = [cb.formato_linea] + [ej.nombre for ej in cb.ejercicios]
            bloques.append((cb.header, lineas))

        # ─── SKILL GIMNÁSTICO ───
        if dia.skill_nombre and dia.skill_gimnastico:
            bloques.append((dia.skill_nombre, [dia.skill_gimnastico]))

        # ─── MUSCULACIÓN (fuerza base + acompañantes) ───
        muscu_lineas = []
        for ej in dia.musculacion:
            muscu_lineas.append(ej.display)
        for ej in dia.acompanantes:
            muscu_lineas.append(ej.display)
        if muscu_lineas:
            bloques.append(("MUSCULACIÓN", muscu_lineas))

        # ─── PLIOMETRÍA ───
        if dia.pliometria:
            p = dia.pliometria
            if p.series > 0:
                lineas = [
                    f"Seated Box Jump {p.series}x{p.reps}",
                    f"Altura: {p.altura}",
                    f"Foco: {p.foco}",
                ]
            else:
                lineas = [
                    "Seated Box Jump — TEST",
                    p.altura,
                    f"Foco: {p.foco}",
                ]
            bloques.append(("PLIOMETRÍA", lineas))

        # ─── ACCESORIOS ───
        if dia.accesorios_lista:
            lineas = [ej.display for ej in dia.accesorios_lista]
            bloques.append(("ACCESORIOS", lineas))

        # ─── SKILL C&J ───
        if dia.skill_cj:
            bloques.append(("SKILL CLEAN & JERK", [dia.skill_cj]))

        # ─── PACING ───
        if dia.pacing:
            p = dia.pacing
            lineas = [
                f"Formato: {p.formato}",
                f"Min Impares: {p.bloque1}",
                f"Min Pares: {p.bloque2}",
                f"Métrica: {p.metrica}",
            ]
            bloques.append(("PACING", lineas))

        # ─── WOD (siempre, vacío para completar) ───
        bloques.append(("WOD", ["[COMPLETAR MANUALMENTE]"]))

        return bloques
